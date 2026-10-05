"""Unit tests for intent detection, clarify sequencing and prompt construction.

No network. The model provider is replaced with a scripted fake, so these
test the pipeline's own logic, not the model.
"""

from __future__ import annotations

import asyncio
import json
import os
import sys
from pathlib import Path
from typing import Any

os.environ["OPENAI_API_KEY"] = ("sk-test_aaaaaaaaaaaaaaaaaaaaaaaaaaaa")
os.environ["API_SHARED_SECRET"] = "unit-test-shared-secret"
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pydantic import ValidationError  # noqa: E402

from config import Settings  # noqa: E402
from core.context import assemble_context, select_relevant_memories  # noqa: E402
from core import events  # noqa: E402
from core.pipeline import clarify, construct, ingest, intent, research, score  # noqa: E402
from core.prompts import system_prompts  # noqa: E402
from core.providers import get_provider  # noqa: E402
from core.providers.base import ModelProvider, ProviderError, ProviderMessage, SearchOutcome, SearchSource  # noqa: E402
from core.prompts.targets import choose_format  # noqa: E402
from main import app  # noqa: E402
from routers.generate import run_pipeline  # noqa: E402
from routers.title import clean_title  # noqa: E402
from schemas.requests import ClarifyQuestion, GenerateRequest, Memory, SourceNote  # noqa: E402

AUTH = {"Authorization": "Bearer unit-test-shared-secret"}


class FakeProvider(ModelProvider):
    """Returns scripted responses in order and records every call."""

    name = "fake"

    def __init__(self, responses: list[Any]) -> None:
        self.responses = list(responses)
        self.calls: list[dict[str, Any]] = []

    def resolve_model(self, requested: str, *, fast: bool = False, step=None) -> str:
        return "fake-fast" if fast else "fake-main"

    def _next(self, **call: Any) -> Any:
        self.calls.append(call)
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response

    async def complete_json(self, *, model, system, messages, max_tokens, temperature, step=""):
        return self._next(kind="json", model=model, system=system, messages=messages)

    async def complete_text(self, *, model, system, messages, max_tokens, temperature, step=""):
        return self._next(kind="text", model=model, system=system, messages=messages)

    async def embed(self, texts, *, step=""):
        """Deterministic bag of words vectors, so relevance is testable without a network."""
        import hashlib
        import re as _re

        vectors = []
        for text in texts:
            vec = [0.0] * 64
            for word in _re.findall(r"[a-z0-9]+", text.lower()):
                if len(word) > 2:
                    vec[int(hashlib.md5(word.encode()).hexdigest(), 16) % 64] += 1.0
            vectors.append(vec)
        return vectors

    async def search(self, *, model, query, step=""):
        return self._next(kind="search", model=model, query=query)

    async def transcribe(self, *, filename, content, content_type):
        return self._next(kind="transcribe", filename=filename, size=len(content))


def make_request(**overrides: Any) -> GenerateRequest:
    base: dict[str, Any] = {"message": "write me a prompt for a support bot"}
    base.update(overrides)
    return GenerateRequest(**base)


def collect(ctx, provider) -> list[Any]:
    async def run() -> list[Any]:
        return [event async for event in run_pipeline(ctx, provider)]

    return asyncio.run(run())


READY_CONSTRUCTION = {
    "content": "## Role\nYou are a support agent for a B2B invoicing product.",
    "stated": ["B2B invoicing product", "Friendly tone"],
    "assumptions": ["English only", "Friendly tone"],
    "structured": {"role": "Support agent", "context": "B2B invoicing", "constraints": ["Be brief"], "format": "Plain text"},
    "readiness": {"status": "ready", "missing": [], "notes": "Complete"},
    "explanation": "Clear role and escalation rules.",
    "memory_candidates": [{"category": "projects", "key": "Product", "value": "B2B invoicing", "confidence": 0.8}],
}

TWO_QUESTIONS = {
    "kind": "prompt_request",
    "sufficient": False,
    "intro": "Happy to build that.",
    "questions": [
        {"text": "What product is it for?"},
        {"text": "What tone should it use?", "options": ["Friendly", "Formal"], "multiple": False},
    ],
}


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

def test_openai_key_is_trimmed_on_load():
    settings = Settings(OPENAI_API_KEY="  sk-test_aaaaaaaaaaaaaaaaaaaaaaaaaaaa\n", API_SHARED_SECRET="a" * 16)
    assert settings.openai_api_key == "sk-test_aaaaaaaaaaaaaaaaaaaaaaaaaaaa"


@pytest.mark.parametrize("bad_key", ["", "   ", "sk-abc\\n trailing junk", "not-a-key", "sk-short"])
def test_malformed_openai_key_fails_loudly(bad_key):
    with pytest.raises(ValidationError):
        Settings(OPENAI_API_KEY=bad_key, API_SHARED_SECRET="a" * 16)


def test_short_shared_secret_fails():
    with pytest.raises(ValidationError):
        Settings(OPENAI_API_KEY="sk-test_aaaaaaaaaaaaaaaaaaaaaaaaaaaa", API_SHARED_SECRET="short")


# ---------------------------------------------------------------------------
# Context assembly
# ---------------------------------------------------------------------------

def test_context_normalizes_blank_profile_fields_and_bounds_history():
    history = [{"role": "user", "content": f"m{i}"} for i in range(30)]
    ctx = assemble_context(make_request(profile={"preferred_name": "  Ada ", "role": "   "}, history=history))
    assert ctx.profile.preferred_name == "Ada"
    assert ctx.profile.role is None
    assert len(ctx.history) == 6
    assert ctx.history[-1].content == "m29"


def test_relevant_memories_keep_about_and_rank_by_overlap():
    memories = [
        Memory(category="about", key="Job", value="Backend engineer"),
        Memory(category="skills", key="Stack", value="Django and Postgres"),
        Memory(category="preferences", key="Coffee", value="Flat white"),
    ]
    selected = select_relevant_memories(memories, "a prompt for reviewing Django migrations")
    assert [m.key for m in selected] == ["Job", "Stack"]


def test_attachments_are_held_unread_until_ingestion_and_earlier_sources_are_restored():
    ctx = assemble_context(
        make_request(attachments=[{"name": "a.png", "type": "image/png", "data_url": "data:image/png;base64,AAA"}])
    )
    assert [a.name for a in ctx.attachments] == ["a.png"]
    assert ctx.sources == []

    state = clarify.start(
        "req", "", [ClarifyQuestion(id="q1", text="Q?")], [SourceNote(kind="attachment", label="brief.pdf", content="x")]
    )
    restored = assemble_context(make_request(clarify_state=state))
    assert [s.label for s in restored.sources] == ["brief.pdf"]


# ---------------------------------------------------------------------------
# Intent detection
# ---------------------------------------------------------------------------

def test_intent_sufficient_request_needs_no_questions():
    decision = intent.parse_intent({"kind": "prompt_request", "sufficient": True, "questions": []}, has_pending_question=False)
    assert decision.kind == "prompt_request"
    assert not decision.needs_clarification


def test_intent_insufficient_request_plans_questions_with_ids():
    decision = intent.parse_intent(TWO_QUESTIONS, has_pending_question=False)
    assert decision.needs_clarification
    assert [q.id for q in decision.questions] == ["q1", "q2"]
    assert decision.questions[0].options is None
    assert decision.questions[1].options == ["Friendly", "Formal"]


def test_intent_caps_questions_at_three_and_drops_duplicates_and_blanks():
    raw = {
        "kind": "prompt_request",
        "sufficient": False,
        "questions": [{"text": "A?"}, {"text": "a?"}, {"text": ""}, {"text": "B?"}, {"text": "C?"}, {"text": "D?"}],
    }
    decision = intent.parse_intent(raw, has_pending_question=False)
    assert [q.text for q in decision.questions] == ["A?", "B?", "C?"]


def test_intent_single_option_falls_back_to_text_answer():
    raw = {"kind": "prompt_request", "sufficient": False, "questions": [{"text": "Which?", "options": ["Only"], "multiple": True}]}
    question = intent.parse_intent(raw, has_pending_question=False).questions[0]
    assert question.options is None
    assert question.multiple is False


def test_intent_insufficient_without_usable_questions_proceeds():
    decision = intent.parse_intent({"kind": "prompt_request", "sufficient": False, "questions": []}, has_pending_question=False)
    assert not decision.needs_clarification


def test_intent_answer_kind_is_only_valid_with_a_pending_question():
    assert intent.parse_intent({"kind": "answer"}, has_pending_question=True).kind == "answer"
    assert intent.parse_intent({"kind": "answer"}, has_pending_question=False).kind == "prompt_request"


def test_intent_conversation_kind():
    assert intent.parse_intent({"kind": "conversation"}, has_pending_question=False).kind == "conversation"


def test_interview_answer_skips_the_model_call():
    state = clarify.start("req", "", [ClarifyQuestion(id="q1", text="Which product?")])
    ctx = assemble_context(make_request(message="Invoicing", clarify_state=state, is_clarify_answer=True))
    provider = FakeProvider([])
    decision = asyncio.run(intent.detect(ctx, provider))
    assert decision.kind == "answer"
    assert provider.calls == []


# ---------------------------------------------------------------------------
# Clarify sequencing
# ---------------------------------------------------------------------------

def _three_question_state():
    return clarify.start(
        "build a bot",
        "Sure.",
        [ClarifyQuestion(id=f"q{i}", text=f"Question {i}?") for i in (1, 2, 3)],
    )


def test_clarify_asks_one_question_at_a_time_in_order():
    state = _three_question_state()
    assert clarify.pending_question(state).id == "q1"
    assert clarify.position(state) == (1, 3)

    state = clarify.record_answer(state, " first ")
    assert clarify.pending_question(state).id == "q2"
    assert clarify.position(state) == (2, 3)
    assert not clarify.is_complete(state)

    state = clarify.record_answer(state, "second")
    state = clarify.record_answer(state, "third")
    assert clarify.is_complete(state)
    assert clarify.question_answer_pairs(state) == [
        ("Question 1?", "first"),
        ("Question 2?", "second"),
        ("Question 3?", "third"),
    ]


def test_clarify_state_is_not_mutated_and_survives_a_json_round_trip():
    state = _three_question_state()
    advanced = clarify.record_answer(state, "first")
    assert state.answers == []
    restored = type(advanced).model_validate(json.loads(advanced.model_dump_json()))
    assert clarify.pending_question(restored).id == "q2"


def test_clarify_rejects_empty_answer_and_answer_after_completion():
    state = clarify.start("req", "", [ClarifyQuestion(id="q1", text="Only?")])
    with pytest.raises(clarify.ClarifyError):
        clarify.record_answer(state, "   ")
    done = clarify.record_answer(state, "yes")
    with pytest.raises(clarify.ClarifyError):
        clarify.record_answer(done, "again")


def test_clarify_caps_planned_questions():
    questions = [ClarifyQuestion(id=f"q{i}", text=f"Q{i}?") for i in range(1, 6)]
    assert len(clarify.start("req", "", questions).questions) == clarify.MAX_QUESTIONS


# ---------------------------------------------------------------------------
# Prompt construction and scoring
# ---------------------------------------------------------------------------

def test_user_turn_restates_original_request_and_every_answer():
    state = clarify.record_answer(clarify.record_answer(_three_question_state(), "alpha"), "beta")
    ctx = assemble_context(make_request(message="beta", clarify_state=state))
    turn = construct.build_user_turn(ctx, state)
    assert "build a bot" in turn
    assert "Question 1?" in turn and "Answer: alpha" in turn
    assert "Question 2?" in turn and "Answer: beta" in turn
    assert "Question 3?" not in turn


def test_user_turn_without_clarification_is_the_message_itself():
    ctx = assemble_context(make_request(message="  summarize this report  "))
    assert construct.build_user_turn(ctx, None) == "summarize this report"


def test_parse_construction_never_lists_an_item_as_both_stated_and_inferred():
    built = construct.parse_construction(READY_CONSTRUCTION)
    assert built.stated == ["B2B invoicing product", "Friendly tone"]
    assert built.assumptions == ["English only"]
    assert built.memory_candidates[0].key == "Product"


def test_parse_construction_without_content_is_an_error_not_a_fabrication():
    from core.providers.base import ProviderResponseError

    with pytest.raises(ProviderResponseError):
        construct.parse_construction({"content": "  ", "assumptions": []})


def test_score_falls_with_assumptions_and_missing_context():
    clean = construct.parse_construction({"content": "x", "stated": ["a"], "assumptions": []})
    loose = construct.parse_construction(
        {
            "content": "x",
            "assumptions": ["a", "b", "c"],
            "readiness": {"status": "needs_context", "missing": ["audience"]},
        }
    )
    assert score.completeness_score(clean) == 90
    assert score.completeness_score(loose) == 90 - 18 - 10 - 15
    assert score.build_diff(loose).inferred == ["a", "b", "c"]


def test_score_is_clamped():
    worst = construct.parse_construction(
        {"content": "x", "assumptions": [str(i) for i in range(20)], "readiness": {"status": "needs_context", "missing": ["a", "b", "c"]}}
    )
    assert score.completeness_score(worst) == score.MIN_SCORE


def test_user_name_inside_a_constructed_prompt_triggers_one_corrective_retry():
    leaky = {**READY_CONSTRUCTION, "content": "## Role - Help Tunde with his project."}
    clean = {**READY_CONSTRUCTION, "content": "## Role - Help the user with their project."}
    provider = FakeProvider([leaky, clean])
    ctx = assemble_context(make_request(profile={"preferred_name": "Tunde"}))
    built = asyncio.run(construct.construct(ctx, None, provider))
    assert "Tunde" not in built.content
    assert len(provider.calls) == 2
    assert "name inside the prompt" in provider.calls[1]["messages"][-1].text


def test_user_name_still_present_after_the_retry_is_scrubbed_not_shipped():
    leaky = {**READY_CONSTRUCTION, "content": "## Role - Help Tunde's team."}
    provider = FakeProvider([leaky, leaky])
    ctx = assemble_context(make_request(profile={"preferred_name": "Tunde"}))
    built = asyncio.run(construct.construct(ctx, None, provider))
    assert "tunde" not in built.content.lower()
    assert len(provider.calls) == 2


# ---------------------------------------------------------------------------
# System prompt
# ---------------------------------------------------------------------------

def test_knowledge_block_lists_empty_fields_as_nothing_stored():
    ctx = assemble_context(make_request(profile={"preferred_name": "Ada"}))
    prompt = system_prompts.answer_system_prompt(ctx)
    assert "Name: Ada" in prompt
    assert "Role: nothing stored" in prompt
    assert "Saved memory (the whole store):\n- nothing stored" in prompt
    assert "say plainly that you have nothing stored" in prompt


def test_partial_memory_is_labelled_as_partial_and_widened_memory_as_the_whole_store():
    from core.context import widen_memory

    memories = [
        {"category": "preferences", "key": "Coffee", "value": "Flat white"},
        {"category": "skills", "key": "Stack", "value": "Django"},
    ]
    ctx = assemble_context(make_request(message="a prompt for Django code review", memories=memories))
    ctx = __import__("dataclasses").replace(ctx, relevant_memories=ctx.all_memories[1:])
    partial = system_prompts.construct_system_prompt(ctx)
    assert "only the items relevant to this message, other items may be stored" in partial
    assert "Django" in partial and "Flat white" not in partial

    whole = system_prompts.answer_system_prompt(widen_memory(ctx))
    assert "Saved memory (the whole store)" in whole
    assert "Flat white" in whole and "Django" in whole


# ---------------------------------------------------------------------------
# Full pipeline, scripted provider
# ---------------------------------------------------------------------------

def test_pipeline_sufficient_request_emits_distinct_stages_then_result():
    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    out = collect(assemble_context(make_request()), provider)
    assert [e.stage for e in out[:-1]] == ["intent", "constructing", "scoring", "finalizing"]
    assert len({e.message for e in out[:-1]}) == 4
    assert out[1].message in events.constructing_pool(0)
    result = out[-1]
    assert result.type == "result"
    assert result.score == 84
    assert result.diff.inferred == ["English only"]


def test_pipeline_underspecified_request_asks_only_the_first_question():
    provider = FakeProvider([TWO_QUESTIONS])
    out = collect(assemble_context(make_request()), provider)
    assert [e.type for e in out] == ["stage", "stage", "clarifying_question"]
    question = out[-1]
    assert question.question.text == "What product is it for?"
    assert (question.position, question.total) == (1, 2)
    assert question.intro == "Happy to build that."
    assert len(provider.calls) == 1


def test_pipeline_full_clarify_sequence_then_result_reflecting_every_answer():
    first = collect(assemble_context(make_request()), FakeProvider([TWO_QUESTIONS]))[-1]

    second = collect(
        assemble_context(make_request(message="Invoicing", clarify_state=first.clarify_state, is_clarify_answer=True)),
        FakeProvider([]),
    )
    assert [e.type for e in second] == ["stage", "stage", "clarifying_question"]
    assert second[-1].question.text == "What tone should it use?"
    assert (second[-1].position, second[-1].total) == (2, 2)
    assert second[-1].intro is None

    provider = FakeProvider([READY_CONSTRUCTION])
    third = collect(
        assemble_context(make_request(message="Friendly", clarify_state=second[-1].clarify_state, is_clarify_answer=True)),
        provider,
    )
    assert third[-1].type == "result"
    assert third[1].message in events.constructing_pool(2)
    final_turn: ProviderMessage = provider.calls[0]["messages"][-1]
    assert "write me a prompt for a support bot" in final_turn.text
    assert "Answer: Invoicing" in final_turn.text
    assert "Answer: Friendly" in final_turn.text


def test_pipeline_direct_question_mid_sequence_keeps_the_pending_question():
    first = collect(assemble_context(make_request()), FakeProvider([TWO_QUESTIONS]))[-1]
    provider = FakeProvider([{"kind": "conversation", "about_user": True}, {"reply": "I have your name stored as Ada. Nothing else.", "memory_candidates": []}])
    ctx = assemble_context(
        make_request(message="what do you know about me?", clarify_state=first.clarify_state, profile={"preferred_name": "Ada"})
    )
    out = collect(ctx, provider)
    assert [e.type for e in out] == ["stage", "stage", "answer"]
    assert out[1].stage == "answering"
    answer = out[-1]
    assert answer.pending_question.text == "What product is it for?"
    assert answer.clarify_state.answers == []
    assert "Role: nothing stored" in provider.calls[1]["system"]


def test_pipeline_typed_reply_classified_as_answer_advances_the_sequence():
    first = collect(assemble_context(make_request()), FakeProvider([TWO_QUESTIONS]))[-1]
    out = collect(
        assemble_context(make_request(message="it is for invoicing", clarify_state=first.clarify_state)),
        FakeProvider([{"kind": "answer"}]),
    )
    assert out[-1].type == "clarifying_question"
    assert out[-1].clarify_state.answers[0].answer == "it is for invoicing"


def test_pipeline_provider_failure_is_an_honest_error_with_no_fabricated_result():
    provider = FakeProvider([ProviderError("boom", status=401, code="invalid_api_key")])
    out = collect(assemble_context(make_request()), provider)
    assert out[-1].type == "error"
    assert out[-1].code == "provider_error"
    assert not any(e.type == "result" for e in out)


# ---------------------------------------------------------------------------
# HTTP surface
# ---------------------------------------------------------------------------

@pytest.fixture
def client():
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_health_is_open_and_everything_else_requires_the_shared_secret(client):
    assert client.get("/health").json()["status"] == "ok"
    assert client.post("/generate", json={"message": "hi"}).status_code == 401
    assert client.post("/title", json={"messages": ["hi"]}, headers={"Authorization": "Bearer wrong"}).status_code == 401
    assert client.post("/transcribe", files={"audio": ("a.webm", b"x", "audio/webm")}).status_code == 401


def test_generate_streams_server_sent_events(client):
    app.dependency_overrides[get_provider] = lambda: FakeProvider([TWO_QUESTIONS])
    response = client.post("/generate", json={"message": "a support bot"}, headers=AUTH)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/event-stream")
    frames = [f for f in response.text.split("\n\n") if f]
    assert [f.split("\n")[0] for f in frames] == ["event: stage", "event: stage", "event: clarifying_question"]
    payload = json.loads(frames[-1].split("\n")[1][len("data: "):])
    assert payload["clarify_state"]["original_message"] == "a support bot"


def test_title_endpoint_cleans_the_inferred_title(client):
    app.dependency_overrides[get_provider] = lambda: FakeProvider([{"title": '"Invoicing support agent."'}])
    response = client.post("/title", json={"messages": ["write me a prompt for a support bot"]}, headers=AUTH)
    assert response.json() == {"title": "Invoicing support agent"}


def test_clean_title_rejects_non_strings():
    assert clean_title(None) == ""


def test_transcribe_endpoint_returns_text(client):
    app.dependency_overrides[get_provider] = lambda: FakeProvider(["hello there"])
    response = client.post("/transcribe", files={"audio": ("a.webm", b"abc", "audio/webm")}, headers=AUTH)
    assert response.json() == {"text": "hello there"}


# ---------------------------------------------------------------------------
# Attachment ingestion
# ---------------------------------------------------------------------------

def make_pdf(text: str) -> bytes:
    """A minimal one page PDF with real extractable text."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import io

    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def b64(data: bytes) -> str:
    import base64

    return base64.b64encode(data).decode()


def test_pdf_text_is_actually_extracted():
    note = ingest.extract_pdf("brief.pdf", make_pdf("Launch date is 12 March for Ledgerly"))
    assert "Launch date is 12 March" in note.content
    assert note.summary.startswith("Read 1 page")
    assert note.kind == "attachment" and note.label == "brief.pdf"


def test_docx_paragraphs_and_tables_are_extracted():
    data = make_docx(["Quarterly goals", "Reduce churn"], [["Metric", "Target"], ["Churn", "3 percent"]])
    note = ingest.extract_docx("goals.docx", data)
    assert "Reduce churn" in note.content
    assert "Churn | 3 percent" in note.content


def test_unreadable_files_become_an_honest_note_not_a_silent_skip():
    attachment = {"name": "broken.pdf", "type": "application/pdf", "content_base64": b64(b"not a pdf")}
    ctx = assemble_context(make_request(attachments=[attachment, {"name": "clip.mp4", "type": "video/mp4"}]))
    notes = asyncio.run(ingest.ingest(ctx.attachments, FakeProvider([]), "m"))
    assert [n.label for n in notes] == ["broken.pdf", "clip.mp4"]
    assert notes[0].summary.startswith("Could not use this file")
    assert "cannot be read yet" in notes[1].summary
    assert all(n.content == "" for n in notes)


def test_long_documents_are_truncated_and_say_so():
    note = ingest.extract_docx("long.docx", make_docx(["word " * 400] * 20))
    assert len(note.content) <= ingest.MAX_CHARS_PER_SOURCE + 20
    assert "only the start was used" in note.summary


def test_images_are_read_with_the_vision_model():
    reply = {"kind": "diagram", "summary": "A three tier architecture diagram.", "content": "Browser to API to Database"}
    provider = FakeProvider([reply])
    note = asyncio.run(ingest.read_image("arch.png", "data:image/png;base64,AAA", provider, "m"))
    assert note.content == "Browser to API to Database"
    assert note.summary == "Read as a diagram. A three tier architecture diagram."
    assert provider.calls[0]["messages"][0].images == ["data:image/png;base64,AAA"]


# ---------------------------------------------------------------------------
# Web research
# ---------------------------------------------------------------------------

def test_search_queries_are_cleaned_deduplicated_and_capped():
    assert research.clean_queries(["a  b", "A B", "", "c", "d"]) == ["a b"]
    assert research.clean_queries("not a list") == []


def test_search_findings_keep_their_sources_and_a_failed_search_is_reported():
    outcome = SearchOutcome(
        summary="FastAPI 0.142 is current. More detail.", sources=[SearchSource("PyPI", "https://pypi.org/p/fastapi")]
    )
    (ok,) = asyncio.run(research.search(["fastapi latest"], FakeProvider([outcome]), "m"))
    assert ok.summary == "FastAPI 0.142 is current."
    assert ok.links[0].url == "https://pypi.org/p/fastapi"
    (failed,) = asyncio.run(research.search(["fastapi latest"], FakeProvider([ProviderError("x")]), "m"))
    assert failed.summary.startswith("The search failed") and failed.content == ""
    # Only one search runs per turn, however many queries were planned.
    provider = FakeProvider([outcome])
    assert len(asyncio.run(research.search(["one", "two"], provider, "m"))) == 1
    assert len(provider.calls) == 1


# ---------------------------------------------------------------------------
# Pipeline with attachments and search
# ---------------------------------------------------------------------------

def test_attachment_is_read_before_intent_and_reflected_in_construction():
    pdf = {"name": "brief.pdf", "type": "application/pdf", "content_base64": b64(make_pdf("Launch date is 12 March for Ledgerly"))}
    built = {**READY_CONSTRUCTION, "researched": ["From brief.pdf: launch date is 12 March"]}
    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, built])
    out = collect(assemble_context(make_request(attachments=[pdf])), provider)

    assert [e.stage for e in out[:-1]] == ["reading", "intent", "constructing", "scoring", "finalizing"]
    assert "Launch date is 12 March" in provider.calls[0]["system"]
    assert "Launch date is 12 March" in provider.calls[1]["system"]
    result = out[-1]
    assert result.diff.researched == ["From brief.pdf: launch date is 12 March"]
    assert [(s.kind, s.label) for s in result.sources] == [("attachment", "brief.pdf")]
    assert "content" not in result.model_dump()["sources"][0]


def test_search_runs_only_when_the_intent_stage_planned_it():
    outcome = SearchOutcome(summary="FastAPI 0.142 is current.", sources=[SearchSource("PyPI", "https://pypi.org/p/fastapi")])
    planned = FakeProvider(
        [{"kind": "prompt_request", "sufficient": True, "search_queries": ["fastapi latest version"]}, outcome, READY_CONSTRUCTION]
    )
    out = collect(assemble_context(make_request()), planned)
    assert [e.stage for e in out[:-1]] == ["intent", "searching", "constructing", "scoring", "finalizing"]
    assert "FastAPI 0.142 is current." in planned.calls[2]["system"]
    assert out[-1].sources[0].kind == "search" and out[-1].sources[0].links[0].title == "PyPI"

    unplanned = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    out = collect(assemble_context(make_request()), unplanned)
    assert "searching" not in [getattr(e, "stage", None) for e in out]
    assert all(c["kind"] != "search" for c in unplanned.calls)


def test_attachment_text_and_early_search_findings_survive_a_clarifying_sequence():
    pdf = {"name": "brief.pdf", "type": "application/pdf", "content_base64": b64(make_pdf("Launch date is 12 March"))}
    plan = {**TWO_QUESTIONS, "search_queries": ["ledgerly competitors"]}
    # Search runs before asking, then the gap survives and the first question is asked.
    first_provider = FakeProvider([plan, SearchOutcome(summary="Competitors include X."), TWO_QUESTIONS])
    first = collect(assemble_context(make_request(attachments=[pdf])), first_provider)[-1]
    assert first.type == "clarifying_question"
    assert {s.kind for s in first.clarify_state.sources} == {"attachment", "search"}
    assert first.clarify_state.search_queries == []

    second = collect(
        assemble_context(make_request(message="Invoicing", clarify_state=first.clarify_state, is_clarify_answer=True)),
        FakeProvider([]),
    )[-1]
    provider = FakeProvider([READY_CONSTRUCTION])
    third = collect(
        assemble_context(make_request(message="Friendly", clarify_state=second.clarify_state, is_clarify_answer=True)),
        provider,
    )
    assert "searching" not in [getattr(e, "stage", None) for e in third]
    assert "Launch date is 12 March" in provider.calls[0]["system"]
    assert "Competitors include X." in provider.calls[0]["system"]
    assert {s.kind for s in third[-1].sources} == {"attachment", "search"}


# ---------------------------------------------------------------------------
# Output format and status variety
# ---------------------------------------------------------------------------

def test_default_output_format_is_plain_and_markdown_only_when_called_for():
    assert choose_format("gpt-4o", "a prompt for a support bot") == "plain"
    assert choose_format("gpt-4o", "write a prompt that outputs a README in Markdown") == "markdown"
    assert choose_format("gpt-4o", "for my Notion workspace") == "markdown"
    assert choose_format("claude-3-5-sonnet", "a prompt for a support bot") == "xml"
    assert choose_format("claude-3-5-sonnet", "output as markdown") == "markdown"


def test_construction_instructions_forbid_markdown_by_default():
    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    out = collect(assemble_context(make_request()), provider)
    assert "No pound signs, no asterisks" in provider.calls[1]["system"]
    assert out[-1].format_style == "plain"

    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    out = collect(assemble_context(make_request(message="a prompt that writes a README in markdown")), provider)
    assert "The destination renders markdown" in provider.calls[1]["system"]
    assert out[-1].format_style == "markdown"


def test_every_stage_pool_mixes_single_words_and_fuller_phrases():
    pools = [
        events.INTENT_POOL, events.INTENT_PENDING_POOL, events.READING_POOL, events.SEARCHING_POOL,
        events.CLARIFYING_FIRST_POOL, events.ANSWERING_POOL, events.SCORING_POOL, events.FINALIZING_POOL,
        events.CONSTRUCTING_BARE_POOL, events.clarifying_pool(2, 3), events.constructing_pool(2),
    ]
    for pool in pools:
        assert any(len(m.split()) == 1 for m in pool), pool
        assert any(len(m.split()) >= 4 for m in pool), pool
        assert len(set(pool)) == len(pool)


def test_stage_messages_always_come_from_their_pool_and_vary_over_calls():
    seen = {events.intent_parsing(False).message for _ in range(80)}
    assert seen <= set(events.INTENT_POOL)
    assert len(seen) > 2
    assert events.finalizing().message in events.FINALIZING_POOL


def test_researched_items_never_duplicate_stated_ones():
    built = construct.parse_construction(
        {
            **READY_CONSTRUCTION,
            "stated": ["Friendly tone"],
            "researched": ["friendly tone", "From a.pdf: 12 March"],
            "assumptions": ["From a.pdf: 12 March", "English only"],
        }
    )
    assert built.researched == ["From a.pdf: 12 March"]
    assert built.assumptions == ["English only"]


def test_a_planned_search_runs_before_asking_and_can_close_the_gap():
    both = {**TWO_QUESTIONS, "search_queries": ["latest python version"]}
    outcome = SearchOutcome(summary="Python 3.14 is current.")
    provider = FakeProvider([both, outcome, {"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    out = collect(assemble_context(make_request()), provider)
    assert [e.stage for e in out[:-1]] == ["intent", "searching", "intent", "constructing", "scoring", "finalizing"]
    assert out[-1].type == "result"
    assert "Python 3.14 is current." in provider.calls[2]["system"]
    assert sum(1 for c in provider.calls if c["kind"] == "search") == 1


def test_a_gap_that_survives_the_search_is_still_asked_and_the_search_is_not_repeated():
    both = {**TWO_QUESTIONS, "search_queries": ["latest python version"]}
    provider = FakeProvider([both, SearchOutcome(summary="Python 3.14 is current."), TWO_QUESTIONS])
    out = collect(assemble_context(make_request()), provider)
    assert out[-1].type == "clarifying_question"
    assert out[-1].clarify_state.search_queries == []
    assert out[-1].clarify_state.sources[0].kind == "search"


def test_time_sensitive_requests_are_detected_and_others_are_not():
    assert research.looks_time_sensitive("a briefing on the latest stable version of Node.js")
    assert research.looks_time_sensitive("what are the current prices and features of the top plans")
    assert not research.looks_time_sensitive("a prompt for a friendly customer support bot")
    assert not research.looks_time_sensitive("summarize the latest email I pasted")


def test_a_time_sensitive_request_gets_a_search_without_a_separate_planning_call():
    provider = FakeProvider(
        [
            {"kind": "prompt_request", "sufficient": True},
            SearchOutcome(summary="Node 26 is current."),
            READY_CONSTRUCTION,
        ]
    )
    message = "a briefing on the latest stable version of Node.js"
    out = collect(assemble_context(make_request(message=message)), provider)
    assert [e.stage for e in out[:-1]] == ["intent", "searching", "constructing", "scoring", "finalizing"]
    assert [c["kind"] for c in provider.calls if c["kind"] != "embed"] == ["json", "search", "json"]
    assert provider.calls[1]["query"] == message
    assert "Node 26 is current." in provider.calls[2]["system"]



def test_provider_tracking_parameter_is_stripped_but_other_query_parameters_stay():
    from core.providers.openai_provider import _strip_tracking

    assert _strip_tracking("https://x.dev/a?utm_source=openai") == "https://x.dev/a"
    assert _strip_tracking("https://x.dev/a?lang=en&utm_source=openai") == "https://x.dev/a?lang=en"
    assert _strip_tracking("https://x.dev/a") == "https://x.dev/a"


# ---------------------------------------------------------------------------
# Memory retrieval, automatic saving and conversation
# ---------------------------------------------------------------------------

STORE = [
    {"category": "skills", "key": "Stack", "value": "React TypeScript Next.js frontend"},
    {"category": "projects", "key": "Product", "value": "Ledgerly invoicing SaaS for finance teams"},
    {"category": "preferences", "key": "Coffee", "value": "Flat white without sugar"},
]


def test_memory_is_consulted_before_intent_and_only_relevant_items_are_selected():
    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    ctx = assemble_context(make_request(message="a prompt for reviewing React TypeScript components", memories=STORE))
    out = collect(ctx, provider)
    assert [m.key for m in out[-1].memory_used] == ["Stack"]
    assert "React TypeScript Next.js frontend" in provider.calls[0]["system"]
    assert "Flat white" not in provider.calls[0]["system"]
    assert "Flat white" not in provider.calls[1]["system"]


def test_a_casual_message_still_consults_memory_before_replying():
    provider = FakeProvider([{"kind": "conversation", "about_user": False}, {"reply": "Databases walk into a bar.", "memory_candidates": []}])
    ctx = assemble_context(make_request(message="tell me a joke about React components", memories=STORE))
    out = collect(ctx, provider)
    assert out[-1].type == "answer"
    assert [m.key for m in out[-1].memory_used] == ["Stack"]
    assert "Flat white" not in provider.calls[1]["system"]


def test_a_message_about_the_user_brings_in_the_whole_store_and_reports_it():
    provider = FakeProvider([{"kind": "conversation", "about_user": True}, {"reply": "You like React and flat whites.", "memory_candidates": []}])
    out = collect(assemble_context(make_request(message="roast me", memories=STORE)), provider)
    assert {m.key for m in out[-1].memory_used} == {"Stack", "Product", "Coffee"}
    assert "Saved memory (the whole store)" in provider.calls[1]["system"]
    assert "Flat white" in provider.calls[1]["system"]


def test_casual_replies_can_propose_memory_and_repeats_are_dropped():
    reply = {
        "reply": "Rust it is.",
        "memory_candidates": [
            {"category": "skills", "key": "Language", "value": "Writes Rust day to day", "confidence": 0.9},
            {"category": "skills", "key": "Stack", "value": "React TypeScript Next.js frontend", "confidence": 0.9},
            {"category": "weird", "key": "Pet", "value": "Has a cat named Miso", "confidence": 0.7},
        ],
    }
    provider = FakeProvider([{"kind": "conversation", "about_user": False}, reply])
    out = collect(assemble_context(make_request(message="btw I write Rust these days", memories=STORE)), provider)
    kept = {(c.key, c.category) for c in out[-1].memory_candidates}
    assert kept == {("Language", "skills"), ("Pet", "other")}


def test_the_conversational_message_travels_apart_from_the_deliverable():
    built = {**READY_CONSTRUCTION, "message": "Here you go, built around your invoicing work."}
    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, built])
    result = collect(assemble_context(make_request()), provider)[-1]
    assert result.message == "Here you go, built around your invoicing work."
    assert "Here you go" not in result.content


def test_a_reply_with_no_text_is_an_error_not_an_empty_message():
    provider = FakeProvider([{"kind": "conversation"}, {"reply": "  ", "memory_candidates": []}])
    out = collect(assemble_context(make_request(message="hello")), provider)
    assert out[-1].type == "error"


def test_memory_selection_falls_back_to_keywords_when_embeddings_fail():
    class NoEmbeddings(FakeProvider):
        async def embed(self, texts, *, step=""):
            raise ProviderError("down")

    provider = NoEmbeddings([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    out = collect(assemble_context(make_request(message="a prompt for reviewing React TypeScript components", memories=STORE)), provider)
    assert out[-1].type == "result"
    assert [m.key for m in out[-1].memory_used] == ["Stack"]


def test_a_follow_up_borrows_the_last_topic_so_a_short_reply_still_finds_memory():
    history = [{"role": "user", "content": "I need a prompt for reviewing React TypeScript components"}]
    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    out = collect(assemble_context(make_request(message="yes do it", memories=STORE, history=history)), provider)
    assert [m.key for m in out[-1].memory_used] == ["Stack"]


def test_em_and_en_dashes_are_replaced_in_every_user_facing_field():
    dash, en = chr(0x2014), chr(0x2013)
    built = construct.parse_construction(
        {
            "message": f"Nice{dash}here you go",
            "content": f"Role{dash}reviewer. Range 1{en}5.",
            "explanation": f"Short{dash}sharp",
        }
    )
    for text in (built.message, built.content, built.explanation):
        assert dash not in text and en not in text
    assert built.content == "Role, reviewer. Range 1-5."

    provider = FakeProvider([{"kind": "conversation"}, {"reply": f"Rust{dash}huh.", "memory_candidates": []}])
    out = collect(assemble_context(make_request(message="hi")), provider)
    assert out[-1].content == "Rust, huh."


# ---------------------------------------------------------------------------
# Usage logging
# ---------------------------------------------------------------------------

def test_cost_estimate_discounts_cached_tokens_and_prices_output_higher():
    from core import usage

    plain = usage.estimate_cost("gpt-4o", 1000, 0, 100)
    cached = usage.estimate_cost("gpt-4o", 1000, 800, 100)
    assert round(plain, 6) == round((1000 * 2.5 + 100 * 10.0) / 1_000_000, 6)
    assert cached < plain
    assert usage.estimate_cost("gpt-4o-mini", 1000, 0, 100) < plain / 10
    # An unknown model is priced as the expensive one, so a mistake overstates spend.
    assert usage.estimate_cost("some-new-model", 1000, 0, 100) == plain


def test_every_call_logs_one_line_and_the_meter_totals_them(caplog):
    import logging

    from core import usage

    caplog.set_level(logging.INFO, logger="data.usage")
    meter = usage.start_meter()
    usage.record("intent", "gpt-4o", prompt_tokens=2000, cached_tokens=1024, completion_tokens=100)
    usage.record("search", "gpt-4o", prompt_tokens=500, completion_tokens=100, extra_cost_usd=usage.SEARCH_CALL_USD)
    meter.log_summary()
    lines = [r.getMessage() for r in caplog.records]
    assert lines[0].startswith("usage step=intent model=gpt-4o prompt_tokens=2000 cached_tokens=1024 completion_tokens=100 est_cost_usd=")
    assert lines[-1].startswith("turn_usage calls=2 prompt_tokens=2500 cached_tokens=1024 completion_tokens=200")
    assert meter.cost_usd > usage.SEARCH_CALL_USD
    assert set(meter.by_step()) == {"intent", "search"}


def test_the_provider_records_real_usage_from_a_chat_response():
    from types import SimpleNamespace

    from core import usage
    from core.providers.openai_provider import OpenAIProvider

    meter = usage.start_meter()
    fake = SimpleNamespace(usage=SimpleNamespace(prompt_tokens=3000, completion_tokens=200, prompt_tokens_details=SimpleNamespace(cached_tokens=2048)))
    OpenAIProvider._record_chat_usage("construct", "gpt-4o", fake)
    rec = meter.records[0]
    assert (rec.step, rec.prompt_tokens, rec.cached_tokens, rec.completion_tokens) == ("construct", 3000, 2048, 200)


def test_a_pipeline_turn_ends_with_a_summary_line(caplog):
    import logging

    caplog.set_level(logging.INFO, logger="data.usage")
    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    collect(assemble_context(make_request()), provider)
    assert any(r.getMessage().startswith("turn_usage calls=") for r in caplog.records)


# ---------------------------------------------------------------------------
# Prompt caching
# ---------------------------------------------------------------------------

def _prefix(system: str) -> str:
    return system.split(system_prompts.REQUEST_CONTEXT_HEADER)[0]


def test_the_static_prefix_is_identical_across_users_and_long_enough_to_be_cached():
    a = assemble_context(make_request(message="a prompt for a bot", profile={"preferred_name": "Ada"}, memories=STORE))
    b = assemble_context(make_request(message="something else entirely", profile={"preferred_name": "Grace", "role": "Chef"}))
    for build in (
        lambda c: system_prompts.intent_system_prompt(c, None),
        system_prompts.construct_system_prompt,
        system_prompts.answer_system_prompt,
    ):
        pa, pb = _prefix(build(a)), _prefix(build(b))
        assert pa == pb
        assert "Ada" not in pa and "Grace" not in pa and "Flat white" not in pa
    # The providers cache prefixes of 1024 tokens or more, about 4 characters each.
    assert len(_prefix(system_prompts.intent_system_prompt(a, None))) > 4096
    assert len(_prefix(system_prompts.construct_system_prompt(a))) > 4096


def test_request_specific_content_comes_after_the_static_prefix():
    ctx = assemble_context(make_request(message="a prompt for reviewing React TypeScript components", profile={"preferred_name": "Ada"}, memories=STORE))
    system = system_prompts.construct_system_prompt(ctx)
    assert system.index("- Name: Ada") > system.index(system_prompts.REQUEST_CONTEXT_HEADER)
    assert system.index("OUTPUT FORMATTING FOR THE PROMPT YOU WRITE") > system.index(system_prompts.REQUEST_CONTEXT_HEADER)



# ---------------------------------------------------------------------------
# Bounded context
# ---------------------------------------------------------------------------

def test_a_long_chat_costs_the_same_as_a_medium_one():
    def tokens(n):
        msgs = [{"role": "user" if i % 2 == 0 else "assistant", "content": ("word " * 200)} for i in range(n)]
        ctx = assemble_context(make_request(history=msgs))
        system = system_prompts.construct_system_prompt(ctx)
        return (len(system) + sum(len(m.content) for m in ctx.history)) // 4

    assert abs(tokens(40) - tokens(25)) < 50
    assert tokens(40) < 6000


def test_older_turns_are_condensed_and_recent_ones_are_clipped():
    from core.context import MAX_DIGEST_CHARS, MAX_LATEST_ASSISTANT_CHARS, MAX_RECENT_CHARS

    msgs = [{"role": "user" if i % 2 == 0 else "assistant", "content": f"turn {i} " + "x" * 3000} for i in range(30)]
    ctx = assemble_context(make_request(history=msgs))
    assert len(ctx.history) == 6
    assert len(ctx.history_digest) <= MAX_DIGEST_CHARS and "turn" in ctx.history_digest
    latest_assistant = [m for m in ctx.history if m.role == "assistant"][-1]
    others = [m for m in ctx.history if m is not latest_assistant]
    assert len(latest_assistant.content) <= MAX_LATEST_ASSISTANT_CHARS + 20
    assert all(len(m.content) <= MAX_RECENT_CHARS + 20 for m in others)
    assert "EARLIER IN THIS CONVERSATION" in system_prompts.intent_system_prompt(ctx, None)


def test_the_latest_prompt_keeps_enough_text_to_be_edited():
    prior = "Role: " + "rule " * 600
    msgs = [{"role": "user", "content": "make a prompt"}, {"role": "assistant", "content": prior}]
    ctx = assemble_context(make_request(message="make it shorter", history=msgs))
    assert len(ctx.history[-1].content) > 2500


def test_intent_sees_a_narrower_window_than_construction():
    msgs = [{"role": "user" if i % 2 == 0 else "assistant", "content": f"m{i}"} for i in range(10)]
    ctx = assemble_context(make_request(history=msgs))
    assert len(intent.build_messages(ctx)) == 4 + 1
    assert len(construct.build_messages(ctx, None)) == 6 + 1


def test_only_the_relevant_few_memories_reach_a_model_prompt():
    many = [{"category": "skills", "key": f"Topic {i}", "value": f"react typescript detail number {i}"} for i in range(60)]
    ctx = __import__("asyncio").run(__import__("core.context", fromlist=["attach_memory"]).attach_memory(
        assemble_context(make_request(message="a prompt for reviewing react typescript code", memories=many)), FakeProvider([])))
    assert len(ctx.all_memories) == 60
    assert len(ctx.relevant_memories) <= 6
    assert system_prompts.construct_system_prompt(ctx).count("react typescript detail number") <= 6


def test_a_query_is_taken_from_the_message_and_bounded():
    assert research.query_from_message("  the latest   Node.js version ") == ["the latest Node.js version"]
    assert research.query_from_message("   ") == []
    assert len(research.query_from_message("x " * 500)[0]) <= research.MAX_QUERY_CHARS


def test_each_step_can_run_on_its_own_model():
    from core.providers.openai_provider import OpenAIProvider

    provider = OpenAIProvider(
        api_key="sk-test_aaaaaaaaaaaaaaaaaaaaaaaaaaaa", default_model="gpt-4o", fast_model="gpt-4o-mini",
        transcribe_model="whisper-1", embed_model="text-embedding-3-small", timeout_seconds=5,
        step_models={"answer": "gpt-4o-mini", "construct": None, "vision": ""},
    )
    assert provider.resolve_model("gpt-4o", step="answer") == "gpt-4o-mini"
    assert provider.resolve_model("gpt-4o", step="construct") == "gpt-4o"
    assert provider.resolve_model("gpt-4o", step="vision") == "gpt-4o"
    assert provider.resolve_model("gpt-4o", step="intent") == "gpt-4o"
    assert provider.resolve_model("gpt-4o", fast=True, step="answer") == "gpt-4o-mini"


# ---------------------------------------------------------------------------
# Limits on the path to the model
# ---------------------------------------------------------------------------

def test_data_rejects_oversized_requests_before_any_model_call(client):
    from schemas.requests import MAX_ATTACHMENTS, MAX_MESSAGE_CHARS

    app.dependency_overrides[get_provider] = lambda: FakeProvider([])
    send = lambda body: client.post("/generate", json=body, headers=AUTH).status_code
    assert send({"message": "x" * (MAX_MESSAGE_CHARS + 1)}) == 422
    assert send({"message": "hi", "history": [{"role": "user", "content": "y"}] * 41}) == 422
    assert send({"message": "hi", "history": [{"role": "user", "content": "y" * 20_001}]}) == 422
    assert send({"message": "hi", "attachments": [{"name": "a.png", "data_url": "d" * 6_000_001}]}) == 422
    assert send({"message": "hi", "attachments": [{"name": "a.pdf", "content_base64": "d" * 14_000_001}]}) == 422
    assert send({"message": "hi", "attachments": [{"name": f"{i}.png"} for i in range(MAX_ATTACHMENTS + 1)]}) == 422
    assert send({"message": "hi", "memories": [{"category": "a", "key": "k", "value": "v" * 601}]}) == 422


def test_a_message_at_the_limit_is_accepted():
    from schemas.requests import MAX_MESSAGE_CHARS

    GenerateRequest(message="x" * MAX_MESSAGE_CHARS)


def test_reference_text_is_bounded_across_all_sources_together():
    big = [SourceNote(kind="attachment", label=f"f{i}.pdf", content="word " * 5000) for i in range(3)]
    block = system_prompts.sources_block(big)
    assert len(block) < system_prompts.MAX_TOTAL_SOURCE_CHARS + 600


def test_output_budgets_leave_room_for_a_thorough_prompt_but_stay_bounded():
    # Enough for a long prompt plus a reasoning model's thinking, and still a ceiling.
    assert 4000 <= construct.CONSTRUCT_MAX_TOKENS <= 10000
    assert ingest.MAX_CHARS_PER_SOURCE <= 8000


def test_lookalike_hyphens_and_spaces_are_normalised_and_a_numeric_range_keeps_its_hyphen():
    text = construct.plain_dashes("non" + chr(0x2011) + "technical, read in 2" + chr(0x2013) + "3 minutes" + chr(0xA0) + "ok")
    assert text == "non-technical, read in 2-3 minutes ok"
    assert construct.plain_dashes("I" + chr(0x2019) + "m " + chr(0x201C) + "sure" + chr(0x201D)) == "I'm \"sure\""


def test_a_failed_step_is_retried_once_on_the_fallback_model():
    class Flaky(FakeProvider):
        def fallback_model(self, model):
            return "steady" if model != "steady" else None

    provider = Flaky([ProviderError("timeout", code="timeout"), {"reply": "ok"}])
    raw = asyncio.run(construct.complete_json_with_fallback(provider, model="big", system="s", messages=[], max_tokens=10, temperature=0.1, step="answer"))
    assert raw == {"reply": "ok"}
    assert [c["model"] for c in provider.calls] == ["big", "steady"]

    lone = FakeProvider([ProviderError("down", status=500)])
    with pytest.raises(ProviderError):
        asyncio.run(construct.complete_json_with_fallback(lone, model="big", system="s", messages=[], max_tokens=10, temperature=0.1, step="answer"))


# ---------------------------------------------------------------------------
# Eval harness, free mock mode
# ---------------------------------------------------------------------------

def _load_eval_module():
    import importlib.util

    spec = importlib.util.spec_from_file_location("run_eval_under_test", Path(__file__).resolve().parent.parent / "evaluation" / "run_eval.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_the_eval_set_runs_in_free_mock_mode_and_every_case_is_well_formed():
    module = _load_eval_module()
    module.MOCK = True
    cases = json.loads(module.EVAL_SET.read_text(encoding="utf-8"))["cases"]
    assert len(cases) >= 40
    results = asyncio.run(module.run_all(cases))
    failed = [(r["id"], r["failures"]) for r in results if not r["passed"]]
    assert failed == []
    assert module.usage.process_meter.cost_usd >= 0


def test_a_live_run_needs_an_explicit_flag_a_confirmation_and_a_scope(monkeypatch, capsys):
    module = _load_eval_module()
    monkeypatch.delenv("EVAL_CONFIRM_SPEND", raising=False)
    run = lambda *argv: (monkeypatch.setattr(sys, "argv", ["run_eval.py", *argv]), module.main())[1]
    assert run("--live", "--full") == 2 and "--yes" in capsys.readouterr().out
    assert run("--live", "--yes") == 2 and "--full" in capsys.readouterr().out
    assert run("--judge") == 2 and "needs --live" in capsys.readouterr().out
    assert run("--save-baseline") == 2


# ---------------------------------------------------------------------------
# Memory summary
# ---------------------------------------------------------------------------

def test_memory_summary_is_written_from_only_the_stored_details(client):
    provider = FakeProvider([{"summary": f"You're Kenny{chr(0x2014)}a developer building a prompt tool."}])
    app.dependency_overrides[get_provider] = lambda: provider
    body = {
        "profile": {"preferred_name": "Kenny", "role": "Developer"},
        "memories": [{"category": "projects", "key": "Current project", "value": "A prompt tool"}],
    }
    response = client.post("/memory-summary", json=body, headers=AUTH)
    assert response.status_code == 200
    assert response.json()["summary"] == "You're Kenny, a developer building a prompt tool."
    sent = provider.calls[0]["messages"][0].text
    assert "- Name: Kenny" in sent and "- [projects] Current project: A prompt tool" in sent
    assert provider.calls[0]["model"] == "fake-fast"
    assert "never use he, she or they" in provider.calls[0]["system"].lower().replace("\n", " ")


def test_memory_summary_with_nothing_stored_says_so_to_the_model_and_failures_are_502(client):
    provider = FakeProvider([{"summary": "Nothing yet."}, ProviderError("down")])
    app.dependency_overrides[get_provider] = lambda: provider
    assert client.post("/memory-summary", json={}, headers=AUTH).status_code == 200
    assert "- nothing stored" in provider.calls[0]["messages"][0].text
    assert client.post("/memory-summary", json={}, headers=AUTH).status_code == 502
    assert client.post("/memory-summary", json={}).status_code == 401


def test_every_terminal_event_carries_what_the_turn_cost():
    from core import usage

    class Metered(FakeProvider):
        async def complete_json(self, *, model, system, messages, max_tokens, temperature, step=""):
            usage.record(step, "gpt-4o", prompt_tokens=1000, cached_tokens=500, completion_tokens=100)
            return await super().complete_json(model=model, system=system, messages=messages, max_tokens=max_tokens, temperature=temperature, step=step)

    result = collect(assemble_context(make_request()), Metered([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION]))[-1]
    assert (result.usage.calls, result.usage.prompt_tokens, result.usage.completion_tokens, result.usage.cached_tokens) == (2, 2000, 200, 1000)
    assert result.usage.est_cost_usd > 0

    question = collect(assemble_context(make_request()), Metered([TWO_QUESTIONS]))[-1]
    assert question.usage.calls == 1

    answer = collect(assemble_context(make_request(message="hello")), Metered([{"kind": "conversation"}, {"reply": "hi", "memory_candidates": []}]))[-1]
    assert answer.usage.calls == 2


# ---------------------------------------------------------------------------
# Feedback review loop
# ---------------------------------------------------------------------------

def _load_feedback_module():
    import importlib.util

    spec = importlib.util.spec_from_file_location("feedback_review_under_test", Path(__file__).resolve().parent.parent / "evaluation" / "feedback_review.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


FEEDBACK_ROWS = [
    {"content": "Role\nOne", "created_at": "2026-10-01T10:00:00Z", "metadata": {"mode": "ready", "score": 80, "templateVersion": "v3", "feedback": {"rating": "up", "request": "a support bot", "score": 80, "templateVersion": "v3", "at": "2026-10-01T10:01:00Z"}}},
    {"content": "Role\nTwo generic text", "created_at": "2026-10-01T11:00:00Z", "metadata": {"mode": "ready", "score": 60, "feedback": {"rating": "down", "request": "A Prompt For Sales", "score": 60, "templateVersion": "v4", "at": "2026-10-01T11:01:00Z"}}},
    {"content": "dup", "created_at": "2026-10-01T12:00:00Z", "metadata": {"mode": "ready", "feedback": {"rating": "down", "request": "a prompt for sales", "score": None, "templateVersion": "v4"}}},
    {"content": "lol", "created_at": "2026-10-01T13:00:00Z", "metadata": {"mode": "answer", "feedback": {"rating": "down", "request": "roast me", "score": None, "templateVersion": "v4"}}},
    {"content": "ignored, malformed", "metadata": {"feedback": {"rating": "maybe"}}},
    {"content": "ignored, no feedback", "metadata": {"mode": "ready"}},
    {"content": "ignored, no metadata", "metadata": None},
]


def test_feedback_rows_are_parsed_and_malformed_ones_skipped():
    module = _load_feedback_module()
    items = module.parse_rows(FEEDBACK_ROWS)
    assert [i.rating for i in items] == ["up", "down", "down", "down"]
    assert items[0].template_version == "v3" and items[1].score == 60


def test_the_summary_breaks_ratings_down_by_template_version():
    module = _load_feedback_module()
    summary = module.summarize(module.parse_rows(FEEDBACK_ROWS))
    assert summary["overall"]["total"] == 4 and summary["overall"]["helpful_rate"] == 25
    assert summary["by_template_version"]["v3"]["helpful_rate"] == 100
    assert summary["by_template_version"]["v4"]["down"] == 3
    assert summary["overall"]["mean_score_when_down"] == 60.0


def test_thumbs_down_items_become_reviewable_drafts_never_finished_cases():
    module = _load_feedback_module()
    cases = module.to_candidate_cases(module.parse_rows(FEEDBACK_ROWS))
    assert [c["input"] for c in cases] == ["A Prompt For Sales", "roast me"]
    assert all(c["needs_review"] and c["id"].startswith("feedback-") for c in cases)
    assert cases[0]["expect"]["decision"] == "ready" and cases[1]["expect"]["decision"] == "conversation"
    assert "DRAFT" in cases[0]["good_prompt_notes"] and "Two generic text" in cases[0]["good_prompt_notes"]


# ---------------------------------------------------------------------------
# Pattern retrieval
# ---------------------------------------------------------------------------

def test_the_pattern_library_is_valid_and_every_pattern_is_substantial():
    from core import patterns as lib

    library = lib.load_library()
    assert len(library) >= 20
    assert len({p.id for p in library}) == len(library)
    for p in library:
        assert len(p.craft) >= 3 and all(len(point) > 40 for point in p.craft), p.id
        assert "\u2014" not in json.dumps(p.craft) and p.summary and p.tags


def test_retrieval_ranks_by_meaning_and_returns_nothing_for_an_unrelated_request():
    from core import patterns as lib

    provider = FakeProvider([])
    found = asyncio.run(lib.retrieve("a customer support agent in chat that handles billing questions and hands off angry customers to a human", provider))
    assert found and found[0].pattern.id == "support-agent"
    assert len(found) <= lib.MAX_PATTERNS
    assert asyncio.run(lib.retrieve("ok go", provider)) == []
    assert asyncio.run(lib.retrieve("   ", provider)) == []


def test_a_failed_embedding_call_means_no_patterns_not_a_failed_turn():
    from core import patterns as lib

    class Down(FakeProvider):
        async def embed(self, texts, *, step=""):
            raise ProviderError("down")

    assert asyncio.run(lib.retrieve("a customer support agent", Down([]))) == []


def test_retrieved_patterns_reach_the_construction_prompt_after_the_static_prefix_and_are_reported():
    provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
    message = "a customer support agent in chat that handles billing questions and hands off angry customers to a human"
    out = collect(assemble_context(make_request(message=message)), provider)
    system = provider.calls[1]["system"]
    assert "KNOWN EFFECTIVE PATTERNS RETRIEVED FOR THIS REQUEST" in system
    assert system.index("KNOWN EFFECTIVE PATTERNS RETRIEVED") > system.index(system_prompts.REQUEST_CONTEXT_HEADER)
    assert "Define the handoff in observable terms" in system
    assert out[-1].patterns_used and out[-1].patterns_used[0] == "Customer support agent"
    assert "patterns" not in provider.calls[0]["system"].lower().split("known effective")[0][-5:]


def test_patterns_can_be_switched_off_for_measurement(monkeypatch):
    from config import get_settings

    monkeypatch.setenv("PATTERNS_ENABLED", "false")
    get_settings.cache_clear()
    try:
        provider = FakeProvider([{"kind": "prompt_request", "sufficient": True}, READY_CONSTRUCTION])
        out = collect(assemble_context(make_request(message="a customer support agent in chat that handles billing questions")), provider)
        assert "KNOWN EFFECTIVE PATTERNS RETRIEVED FOR THIS REQUEST" not in provider.calls[1]["system"]
        assert out[-1].patterns_used == []
    finally:
        monkeypatch.delenv("PATTERNS_ENABLED")
        get_settings.cache_clear()


def test_a_clarified_request_retrieves_on_the_request_and_the_answers_together():
    state = clarify.record_answer(clarify.start("write a prompt for a bot", "", [ClarifyQuestion(id="q1", text="What kind?")]), "customer support for billing questions")
    ctx = assemble_context(make_request(message="customer support for billing questions", clarify_state=state, is_clarify_answer=True))
    assert construct.pattern_query(ctx, state) == "write a prompt for a bot customer support for billing questions"


# ---------------------------------------------------------------------------
# Questions about the outside world are searched, not recalled
# ---------------------------------------------------------------------------

def test_a_question_about_a_real_person_is_looked_up_in_its_own_stage_and_answered_from_the_findings():
    findings = SearchOutcome(
        summary="Lila Ibrahim is Chief Operating Officer of Google DeepMind since 2018.",
        sources=[SearchSource("Axios", "https://www.axios.com/2018/04/11/exclusive-deepmind-hires-lila-ibrahim-as-first-coo")],
    )
    provider = FakeProvider(
        [
            {"kind": "conversation", "lookup": True, "search_queries": ["Lila Ibrahim biography"], "about_user": False},
            findings,
            {"reply": "Lila Ibrahim is the COO of Google DeepMind."},
        ]
    )
    out = collect(assemble_context(make_request(message="who is Lila Ibrahim")), provider)

    assert [e.stage for e in out[:-1]] == ["intent", "lookup", "answering"]
    assert [c["kind"] for c in provider.calls if c["kind"] != "embed"] == ["json", "search", "json"]
    # The findings reach the model that writes the answer, as reference material.
    assert "Chief Operating Officer of Google DeepMind" in provider.calls[2]["system"]
    answer = out[-1]
    assert answer.type == "answer"
    assert [(s.kind, s.label) for s in answer.sources] == [("search", "Lila Ibrahim biography")]
    assert answer.sources[0].links[0].title == "Axios"


def test_a_plain_who_is_a_named_person_is_searched_even_when_the_intent_model_forgot():
    provider = FakeProvider(
        [
            {"kind": "conversation", "about_user": False},
            SearchOutcome(summary="Ada Lovelace was a mathematician."),
            {"reply": "Ada Lovelace was a mathematician."},
        ]
    )
    out = collect(assemble_context(make_request(message="Who was Ada Lovelace?")), provider)
    assert [e.stage for e in out[:-1]] == ["intent", "lookup", "answering"]
    assert provider.calls[1]["query"] == "Ada Lovelace"


def test_questions_about_the_user_or_general_knowledge_are_never_searched():
    for message, raw in [
        ("who am I?", {"kind": "conversation", "about_user": True}),
        ("who is my manager", {"kind": "conversation", "about_user": True}),
        ("how does a b-tree work", {"kind": "conversation", "about_user": False}),
        ("tell me about recursion", {"kind": "conversation", "about_user": False}),
    ]:
        provider = FakeProvider([raw, {"reply": "ok"}])
        out = collect(assemble_context(make_request(message=message)), provider)
        assert "lookup" not in [getattr(e, "stage", None) for e in out], message
        assert all(c["kind"] != "search" for c in provider.calls), message
        assert out[-1].sources == []


def test_a_name_that_is_the_users_own_or_already_stored_is_answered_from_memory():
    provider = FakeProvider([{"kind": "conversation", "about_user": False}, {"reply": "That is you."}])
    request = make_request(message="who is Kenny", profile={"preferred_name": "Kenny"})
    out = collect(assemble_context(request), provider)
    assert all(c["kind"] != "search" for c in provider.calls)
    assert "lookup" not in [getattr(e, "stage", None) for e in out]


def test_a_failed_lookup_is_reported_to_the_answering_model_not_papered_over():
    provider = FakeProvider(
        [
            {"kind": "conversation", "lookup": True, "search_queries": ["Some Person"]},
            ProviderError("down", status=500),
            {"reply": "I could not look that up right now."},
        ]
    )
    out = collect(assemble_context(make_request(message="who is Some Person")), provider)
    assert "The search failed" in provider.calls[2]["system"]
    assert out[-1].type == "answer"


def test_lookup_subject_only_fires_on_a_named_subject():
    from core.pipeline.research import lookup_subject

    assert lookup_subject("who is Lila Ibrahim") == "Lila Ibrahim"
    assert lookup_subject("Tell me about Chinua Achebe") == "Chinua Achebe"
    for plain in ["who am I", "who is my boss", "tell me about recursion", "who is responsible for this bug", "who is it"]:
        assert lookup_subject(plain) == "", plain


# ---------------------------------------------------------------------------
# Memory, search and conversation are inputs that combine
# ---------------------------------------------------------------------------

MEMORIES = [
    {"category": "projects", "key": "Current project", "value": "A CLI tool that parses server logs, written in Rust"},
    {"category": "skills", "key": "Deployment", "value": "Ships a static musl binary to Alpine containers"},
]


def test_a_lookup_answer_is_written_from_the_findings_and_the_users_stored_context_together():
    provider = FakeProvider(
        [
            {"kind": "conversation", "lookup": True, "search_queries": ["latest stable Rust release"], "about_user": True},
            SearchOutcome(summary="Rust 1.99.0 is the latest stable release.", sources=[SearchSource("Rust", "https://rust-lang.org")]),
            {"reply": "Rust 1.99.0 is current, and it suits your log parser."},
        ]
    )
    request = make_request(message="What is the latest stable Rust release, and does it matter for what I'm building?", memories=MEMORIES)
    out = collect(assemble_context(request), provider)

    system = provider.calls[-1]["system"]
    assert "Rust 1.99.0 is the latest stable release." in system
    assert "A CLI tool that parses server logs" in system and "static musl binary" in system
    answer = out[-1]
    assert {m.key for m in answer.memory_used} == {"Current project", "Deployment"}
    assert [s.kind for s in answer.sources] == ["search"]
    assert "TODAY'S DATE" in system


def test_searching_keeps_the_best_stored_context_even_when_the_words_do_not_match():
    provider = FakeProvider(
        [
            {"kind": "conversation", "lookup": True, "search_queries": ["who is Ada Lovelace"]},
            SearchOutcome(summary="Ada Lovelace wrote the first program."),
            {"reply": "ok"},
        ]
    )
    request = make_request(message="Tell me about Ada Lovelace for a talk", memories=MEMORIES)
    out = collect(assemble_context(request), provider)
    assert len(out[-1].memory_used) >= 1


def test_a_prompt_built_after_a_search_still_sees_stored_memory():
    provider = FakeProvider(
        [
            {"kind": "prompt_request", "sufficient": True, "search_queries": ["fastapi latest version"]},
            SearchOutcome(summary="FastAPI 0.142 is current."),
            READY_CONSTRUCTION,
        ]
    )
    request = make_request(message="Write a prompt for upgrading my service to the latest FastAPI", memories=MEMORIES)
    out = collect(assemble_context(request), provider)
    system = provider.calls[-1]["system"]
    assert "FastAPI 0.142 is current." in system
    assert "A CLI tool that parses server logs" in system
    assert out[-1].memory_used
